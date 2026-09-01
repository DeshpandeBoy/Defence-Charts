from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path('/Users/dhanyarao/Downloads/Integration Assessment.docx')
OUTPUT = Path('/Users/dhanyarao/Documents/Defence/deliverables/Integration Assessment - Professional.docx')
WORK = Path('/private/tmp/integration-assessment-assets')

NAVY = '163A5F'
BLUE = '2E74B5'
TEXT = '1F2937'
MUTED = '5B6573'
PALE_BLUE = 'EAF1F8'
LIGHT_BLUE = 'F5F8FC'
RULE = 'B8C7D9'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(width_dxa))
    tc_w.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_layout = tbl_pr.first_child_found_in('w:tblLayout')
    if tbl_layout is None:
        tbl_layout = OxmlElement('w:tblLayout')
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn('w:type'), 'fixed')
    tbl_ind = tbl_pr.first_child_found_in('w:tblInd')
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    tr_pr.append(header)


def set_font(run, size=11, color=TEXT, bold=None, italic=None):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_rule(paragraph, color=RULE, size='8', space='6'):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.first_child_found_in('w:pBdr')
    if borders is None:
        borders = OxmlElement('w:pBdr')
        p_pr.append(borders)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), space)
    bottom.set(qn('w:color'), color)
    borders.append(bottom)


def set_keep(paragraph, with_next=False):
    p_pr = paragraph._p.get_or_add_pPr()
    tag = 'w:keepNext' if with_next else 'w:keepLines'
    if p_pr.find(qn(tag)) is None:
        p_pr.append(OxmlElement(tag))


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    set_keep(p)
    set_font(p.add_run(text), 10.8)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.left_indent = Inches(0.31)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    set_keep(p)
    set_font(p.add_run(text), 10.8)
    return p


def add_heading(doc, text, page_break_before=False):
    p = doc.add_paragraph()
    p.paragraph_format.page_break_before = page_break_before
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(7)
    set_keep(p, with_next=True)
    set_font(p.add_run(text), size=15, color=NAVY, bold=True)
    return p


def add_image(doc, path, width, alt_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=width)
    doc.inline_shapes[-1]._inline.docPr.set('descr', alt_text)
    return p


def set_styles(doc):
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name in ['List Bullet', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        style.font.size = Pt(10.8)


def setup_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(2)
    set_font(header.add_run('INTEGRATION ASSESSMENT  |  TECHNICAL REVIEW'), 8.5, MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(2)
    set_font(footer.add_run('Integration Assessment  •  '), 8.5, MUTED)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run('INTEGRATION ASSESSMENT'), 22, NAVY, bold=True)
    add_rule(p, NAVY, '14', '7')
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(13)
    set_font(subtitle.add_run('Insight Summary to Patient Hub Dashboard'), 12, MUTED, italic=True)
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    set_font(meta.add_run('Scope: current-state data flows, integration alignment, and technical considerations.'), 9.5, MUTED)


def build():
    WORK.mkdir(parents=True, exist_ok=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(SOURCE) as archive:
        for image in ('image1.png', 'image2.png'):
            (WORK / image).write_bytes(archive.read(f'word/media/{image}'))

    doc = Document()
    setup_page(doc)
    set_styles(doc)
    add_title_block(doc)

    add_heading(doc, '1. Current Insight Summary Flow')
    for item in [
        'The current Insight Summary uses synthetic patient-therapy data as its source.',
        'Day-level summary data is organized in JSON format.',
        'Existing session values, including AHI and usage duration, are used to derive additional values from the available parameters.',
        'The pre-computed values are supplied to the LLM as context for generation of the Patient Summary.',
    ]:
        add_bullet(doc, item)
    add_image(doc, WORK / 'image1.png', Inches(6.65), 'Current Insight Summary processing and summary-generation flow diagram.')

    doc.add_page_break()
    add_heading(doc, '2. Current Patient Hub Flow')
    for item in [
        'Patient therapy and ventilator data from modem and SD-card sources is aggregated through the aggregation service and stored in DocDB.',
        'DocDB (an Amazon database) stores day-level summary data in JSON format.',
        'DSD parameters are transformed from JSON into structured Data Mart tables. This data is used to calculate required values through microservices and to populate the schema tables required for Patient Hub fields.',
        'Data Mart tables are created according to the data required for the Patient Hub Dashboard.',
        'The Data Mart is updated whenever new source data is received.',
        'Cumulative calculated data for predefined periods is computed and stored in PostgreSQL tables.',
    ]:
        add_bullet(doc, item)
    add_image(doc, WORK / 'image2.png', Inches(6.65), 'Current Patient Hub aggregation and data-population flow diagram.')

    doc.add_page_break()
    add_heading(doc, '3. Current-State Comparison')
    caption = doc.add_paragraph()
    caption.paragraph_format.space_after = Pt(5)
    set_font(caption.add_run('Comparison of the current Insight Summary implementation and the required Patient Hub adaptation.'), 9.5, MUTED, italic=True)
    data = [
        ['SL.', 'Details', 'Present State', 'Need to Adapt'],
        ['1', 'Raw Data Source', 'Synthetic Patient Therapy Data', 'Patient Therapy and Vent Data from Modems/SD Cards'],
        ['2', 'Raw Data Representation', 'Restructured list of arrays (patient ID)', 'Combined day-wise summary data (DSD) in DocDB'],
        ['3', 'Input Data Source', 'Synthetic Data', 'DocDB for combined session data; PostgreSQL for aggregated dashboard-compatible data'],
        ['4', 'Computations', 'Existing parameters (AHI, usage duration, etc.) plus derived computations (average usage, mask-fit %, etc.) in Insight Summary code', 'Available pre-computed values on a microservice, reusable through an API'],
        ['5', 'Past Data Availability', 'No threshold set', 'Maximum two years of historical data'],
    ]
    table = doc.add_table(rows=0, cols=4)
    table.style = 'Table Grid'
    widths = [480, 1920, 3320, 3520]
    set_table_geometry(table, widths)
    for ridx, row in enumerate(data):
        cells = table.add_row().cells
        for cidx, text in enumerate(row):
            cell = cells[cidx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ridx == 0:
                set_cell_shading(cell, NAVY)
            elif ridx % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if cidx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_font(run, 9.2 if ridx else 9.4, 'FFFFFF' if ridx == 0 else TEXT, bold=(ridx == 0))
    mark_header_row(table.rows[0])

    add_heading(doc, '4. Integration Considerations', page_break_before=True)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(9)
    set_font(intro.add_run('Integration of the patient-therapy data summary feature with Care Orchestrator requires the following technical decisions and validation activities.'), 10.8)
    for item in [
        'Define the scope for Sleep and Vent data: determine whether both data types must be supported, whether they are unique, and whether the feature should accommodate one or both.',
        'DocDB contains Sleep and Vent session data, with the `_vent` parameter distinguishing aggregated ventilator data from Sleep data. This JSON-based representation is consistent with the current Insight Summary flow.',
        'The current Insight Summary prompt uses computed values to generate summaries for DME and Clinician personas. If Vent data is included, the computations and prompt must be redesigned to incorporate the relevant ventilator calculations.',
        'PostgreSQL contains general computed values, including AHI, Avg_Usage, and leak percentage, which can be reused through microservice functions. However, derived computations used for the Clinician-persona summary are not currently available in PostgreSQL.',
        'Existing computations can be reused without change, with additional computations implemented as extensions. Alternatively, the additional computations can be maintained separately outside the AI layer.',
        'Data-access contract details and API specifications require verification.',
    ]:
        add_numbered(doc, item)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
